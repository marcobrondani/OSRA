"""The workbooks: generation from the pack, export, import and the round trip
(ADR-0008, TR-60 to TR-63, TR-103)."""

from __future__ import annotations

import io
import itertools
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

from com.brondani.osra import workbook
from com.brondani.osra.canonical import yaml_text
from com.brondani.osra.store import Actor, Store
from com.brondani.osra.validate import check_assessment

from .conftest import EXAMPLE, REPO
from .test_reports import SUMMARY

HUMAN = Actor(author="Assessor")
KINDS = ("assessment", "substrate", "failures", "trust", "scoring", "summary")


@pytest.fixture
def scored(tmp_path, pack):
    counter = itertools.count()

    def clock():
        t = next(counter)
        return f"2026-10-01T10:{t // 60 % 60:02d}:{t % 60:02d}Z"

    root = tmp_path / "a"
    shutil.copytree(EXAMPLE, root)
    store = Store(root, pack, clock=clock)
    store.record(HUMAN)
    store.rate(HUMAN, "DEP-02", {n: {"score": 3} for n in workbook.ENTERED_FACTORS})
    store.summarise(HUMAN, "DEP-01", SUMMARY)
    store.resolve_tie(HUMAN, ["DEP-02", "DEP-01"], "Recorded for the test.")
    for register in ("substrate", "failures", "trust", "scoring", "summary"):
        store.confirm(HUMAN, register, interactive=True)
    store.score(HUMAN)
    return store


def canonical(pack, documents):
    return {k: yaml_text(documents[k], f"{k}.schema.json", pack.schemas) for k in documents}


def test_export_then_import_is_the_identity(scored, pack, tmp_path):
    docs = scored.documents()
    paths = workbook.export(docs, scored.results(), pack, tmp_path / "wb")
    assert [p.name for p in paths] == list(workbook.FILES.values())
    imported = workbook.import_workbooks(paths, pack)
    assert imported.revised and imported.notes == []
    assert canonical(pack, imported.documents) == canonical(pack, {k: docs[k] for k in KINDS})


def test_round_trip_through_the_store(scored, pack, tmp_path):
    paths = workbook.export(scored.documents(), scored.results(), pack, tmp_path / "wb")
    target = Store(tmp_path / "b", pack, clock=lambda: "2026-10-02T09:00:00Z")
    target.import_documents(Actor(author="Assessor", surface="import"), workbook.import_workbooks(paths, pack).documents,
                            source=[p.name for p in paths], attribute=False)
    assert canonical(pack, target.documents()) == canonical(pack, scored.documents())
    for kind in KINDS[1:]:  # the registers were written by the store, so the files match byte for byte
        assert (target.root / f"{kind}.yaml").read_bytes() == (scored.root / f"{kind}.yaml").read_bytes(), kind
    assert target.integrity() == []


def test_generated_data_round_trips(pack, tmp_path):
    """TR-103: the round trip holds for generated data, not only the example."""
    from com.brondani.osra.verify import build_registers
    from com.brondani.osra.fixtures import load_fixtures

    fixtures, _ = load_fixtures(pack)
    for scenario in fixtures.scenarios.values():
        docs = build_registers(scenario, pack)
        docs["assessment"] = {"schema_version": 1, "kind": "assessment", "method_pack": {"id": "osra", "version": "1.2"},
                              "system": {"name": scenario["name"], "owner": "o", "regulatory_classification": "c",
                                         "boundary": "b"},
                              "assessment_type": "full", "deployment_mode": {"declared": "C"},
                              "agent_access": "read-only", "status": "active"}
        out = tmp_path / scenario["id"]
        imported = workbook.import_workbooks(workbook.export(docs, None, pack, out), pack)
        assert canonical(pack, {k: imported.documents[k] for k in docs}) == canonical(pack, docs), scenario["id"]


def test_computed_columns_hold_engine_values_and_no_cell_holds_a_formula(scored, pack, tmp_path):
    paths = workbook.export(scored.documents(), scored.results(), pack, tmp_path / "wb")
    for path in paths:
        wb = openpyxl.load_workbook(path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", (path.name, ws.title, cell.coordinate)
    ws = openpyxl.load_workbook(paths[1])["Failure Surface Register"]
    headers = [c.value for c in ws[3]]
    severity = headers.index("Severity (computed)") + 1
    assert ws.cell(row=6, column=severity).value == "Low"  # FM-03 after its tested fallback
    assert ws.cell(row=6, column=severity).protection.locked
    assert not ws.cell(row=6, column=1).protection.locked
    assert ws.protection.sheet
    scoring = openpyxl.load_workbook(paths[3])["Convergence Scoring"]
    headers = [c.value for c in scoring[3]]
    assert scoring.cell(row=4, column=headers.index("Score (computed)") + 1).value == "33.0"


def test_validation_lists_come_from_the_pack(pack, tmp_path):
    paths = workbook.export({}, None, pack, tmp_path / "t")
    lists = openpyxl.load_workbook(paths[0])["Lists"]
    columns = {lists.cell(row=1, column=c).value: [lists.cell(row=r, column=c).value for r in range(2, lists.max_row + 1)]
               for c in range(1, lists.max_column + 1)}
    vocab = workbook.Vocabulary(pack)
    assert [v for v in columns["layer"] if v] == vocab.labels("layer")
    assert [v for v in columns["visibility"] if v] == ["Visible", "Known-Unmonitored", "Invisible"]


def test_templates_are_deterministic(pack, tmp_path):
    first = workbook.export({}, None, pack, tmp_path / "one")
    second = workbook.export({}, None, pack, tmp_path / "two")
    for a, b in zip(first, second):
        assert a.read_bytes() == b.read_bytes(), a.name
    with zipfile.ZipFile(first[0]) as archive:
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist())


def test_text_that_looks_like_a_formula_stays_text(scored, pack, tmp_path):
    scored.set(HUMAN, "FM-02", {"description": '=HYPERLINK("http://example.invalid","x")'})
    paths = workbook.export(scored.documents(), None, pack, tmp_path / "wb")
    cell = openpyxl.load_workbook(paths[1])["Failure Surface Register"]["D5"]
    assert cell.data_type == "s"
    imported = workbook.import_workbooks(paths, pack)
    assert imported.documents["failures"]["failure_modes"][1]["description"].startswith("=HYPERLINK")


def test_a_confirmed_register_edited_in_the_workbook_returns_to_draft(scored, pack, tmp_path):
    paths = workbook.export(scored.documents(), scored.results(), pack, tmp_path / "wb")
    wb = openpyxl.load_workbook(paths[0])
    wb["Dependency Chain"]["C4"] = "Hosted base model, renamed"
    wb.save(paths[0])
    imported = workbook.import_workbooks(paths, pack)
    assert imported.documents["substrate"]["state"] == "draft"
    assert "confirmation" not in imported.documents["substrate"]
    assert [(n.code, n.entity) for n in imported.notes] == [("OSRA-W622", "substrate")]
    assert imported.documents["failures"]["state"] == "confirmed"


def test_a_value_typed_by_hand_is_read_by_name_or_label(scored, pack, tmp_path):
    paths = workbook.export(scored.documents(), None, pack, tmp_path / "wb")
    wb = openpyxl.load_workbook(paths[1])
    wb["Failure Surface Register"]["F4"] = "  NEVER "
    wb["Failure Surface Register"]["C4"] = "silent failure"
    wb.save(paths[1])
    fm = workbook.import_workbooks(paths, pack).documents["failures"]["failure_modes"][0]
    assert (fm["detection"]["latency"], fm["type"]) == ("never", "silent")


# -- the workbooks as published in v1.2 --------------------------------------------


def published_copy(tmp_path) -> list[Path]:
    out = tmp_path / "published"
    out.mkdir()
    for name in workbook.FILES.values():
        shutil.copy(REPO / "templates" / name, out / name)
    return [out / name for name in workbook.FILES.values()]


def test_a_blank_published_template_imports_nothing(pack, tmp_path):
    imported = workbook.import_workbooks(published_copy(tmp_path), pack)
    assert not imported.revised
    assert imported.documents["substrate"]["dependencies"] == []
    assert imported.documents["failures"]["failure_modes"] == []
    assert imported.documents["scoring"]["scores"] == []
    assert all(n.code == "OSRA-W622" for n in imported.notes)


def test_a_filled_published_template_imports_as_a_draft_with_its_gaps(pack, tmp_path):
    paths = published_copy(tmp_path)
    wb = openpyxl.load_workbook(paths[0])
    for row in wb["System Boundary"].iter_rows():
        values = {"AI System Name": "Screening", "System Owner": "Risk", "Regulatory Classification": "DORA"}
        if row[0].value in values:
            row[2].value = values[row[0].value]
    chain = wb["Dependency Chain"]
    chain["D4"], chain["F4"], chain["G4"], chain["H4"], chain["I4"] = "Model provider", "Y", "Known-Unmonitored", "N", "N"
    wb.save(paths[0])
    wb = openpyxl.load_workbook(paths[1])
    sheet = wb["Failure Surface Register"]
    sheet["F4"], sheet["G4"], sheet["H4"], sheet["I4"], sheet["J4"] = "Dashboard", "Never", "Low", "Critical", "N"
    wb.save(paths[1])
    imported = workbook.import_workbooks(paths, pack)
    docs = imported.documents
    docs["assessment"]["system"]["boundary"] = "Model and feed"
    assert [p for p in check_assessment(pack, docs) if p.code.startswith("OSRA-E")] == []
    [dep] = docs["substrate"]["dependencies"]
    assert (dep["id"], dep["owner_type"], dep["single_point"], dep["visibility"]) == ("DEP-01", "unknown", True, "known-unmonitored")
    [fm] = docs["failures"]["failure_modes"]
    assert (fm["type"], fm["impact"], fm["detection"]["confidence"]) == ("silent", "critical", "low")
    assert "severity" not in fm  # computed columns are ignored and recomputed
    assert all(d["state"] == "draft" for k, d in docs.items() if k != "assessment")
    details = " ".join(n.message for n in imported.notes)
    assert "no owner type" in details and "give it with --boundary" in details
    assert "DEP-02: not imported" in details


# -- untrusted input ---------------------------------------------------------------------


def test_an_oversized_archive_is_refused(scored, pack, tmp_path, monkeypatch):
    paths = workbook.export(scored.documents(), None, pack, tmp_path / "wb")
    monkeypatch.setattr(workbook, "MAX_ARCHIVE_BYTES", 1000)
    imported = workbook.import_workbooks(paths[:1], pack)
    assert [n.code for n in imported.notes] == ["OSRA-E302"]
    assert "over the limit" in imported.notes[0].message


def test_an_archive_with_an_unsafe_path_is_refused(pack, tmp_path):
    path = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../escape.xml", "<x/>")
    imported = workbook.import_workbooks([path], pack)
    assert [n.code for n in imported.notes] == ["OSRA-E302"]


def test_a_workbook_from_another_method_version_is_reported(scored, pack, tmp_path):
    paths = workbook.export(scored.documents(), None, pack, tmp_path / "wb")
    wb = openpyxl.load_workbook(paths[2])
    meta = wb["OSRA"]
    for row in meta.iter_rows():
        if row[0].value == "method_pack":
            row[1].value = "osra 1.1"
    wb.save(paths[2])
    imported = workbook.import_workbooks(paths, pack)
    assert ("OSRA-E205", "Phase3_Trust_Surface_Register.xlsx") in [(n.code, n.entity) for n in imported.notes]
