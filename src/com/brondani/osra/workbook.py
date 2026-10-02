"""The OSRA workbooks: generation, export and import (ADR-0008, TR-60 to
TR-63).

The workbooks are generated from the method pack: validation lists, guides
and anchors come from the same source as the engine. They are data-entry
only. No cell holds a formula; every computed value (severity, the silent
failure flag, the horizon score, the trust gap, chain depth, conditions,
category, flag, clock, score, rank) is written by the engine into a
protected cell whose header says it is computed.

The revised workbooks carry every field the model holds, including register
state, retired identifiers, trust chains and, in a column marked not to be
edited, each entry's provenance, so an export followed by an import gives
back the same assessment. Workbooks as published in v1.2 carry none of that;
importing one starts a draft assessment and lists what it could not supply
(FR-70).

Workbook files are untrusted input (TR-86): XML is parsed through
defusedxml, external links are not loaded, an archive larger than the limit
is refused, and text that would read as a formula is written as text.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl
import openpyxl.xml
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .canonical import json_line
from .engine import id_number
from .errors import Diagnostic, diagnostic
from .pack import MethodPack

FORMAT = "osra-workbook"
FORMAT_VERSION = 1
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024  # uncompressed, per workbook
SPARE_ROWS = 100  # rows open for entry below the data
FILES = {
    "phase1": "Phase1_Substrate_Map.xlsx",
    "phase2": "Phase2_Failure_Surface_Register.xlsx",
    "phase3": "Phase3_Trust_Surface_Register.xlsx",
    "phase4": "Phase4_Convergence_Map.xlsx",
}
_FIXED = datetime(2026, 1, 1)
_COMPUTED_FILL = PatternFill("solid", fgColor="E7E9EC")
_HEADER_FONT = Font(bold=True)

if not openpyxl.xml.DEFUSEDXML:  # pragma: no cover - defusedxml is a dependency
    raise ImportError("defusedxml is required to read workbooks safely")


# -- column specifications ------------------------------------------------------
#
# (header, path in the entry, kind). Kinds: id, text, bool, list, json, or
# ("enum", vocabulary). Enumerated values are written with the workbook label
# the method pack gives them and read back from id, name or label.

Column = tuple[str, str, Any]

SUBSTRATE: list[Column] = [
    ("Dependency ID", "id", "id"),
    ("Dependency Layer", "layer", ("enum", "layer")),
    ("Specific Dependency", "name", "text"),
    ("Provider / Owner", "provider", "text"),
    ("Owner Type", "owner_type", ("enum", "owner_type")),
    ("Location / Region", "location", "text"),
    ("Single Point of Dependency? (Y/N)", "single_point", "bool"),
    ("Visibility", "visibility", ("enum", "visibility")),
    ("Fallback Available? (Y/N/Partial)", "fallback", ("enum", "fallback")),
    ("Fallback Tested? (Y/N)", "fallback_tested", "bool"),
    ("Notes / Risks", "notes", "text"),
    ("Provenance (do not edit)", "provenance", "json"),
]
FAILURES: list[Column] = [
    ("Failure ID", "id", "id"),
    ("Dependency ID", "dependency", "text"),
    ("Failure Type", "type", ("enum", "type")),
    ("Failure Description", "description", "text"),
    ("Detection Mechanism", "detection.mechanism", "text"),
    ("Detection Latency", "detection.latency", ("enum", "latency")),
    ("Detection Confidence", "detection.confidence", ("enum", "confidence")),
    ("Impact Level", "impact", ("enum", "impact")),
    ("Tested Fallback In Place? (Y/N)", "tested_fallback", "bool"),
    ("Materialisation Horizon", "materialisation_horizon", ("enum", "materialisation_horizon")),
    ("Propagation Path", "propagation.path", "text"),
    ("Downstream Impact", "propagation.downstream_impact", "text"),
    ("Who Would Notice Downstream", "propagation.downstream_detection", "text"),
    ("Notes", "notes", "text"),
    ("Provenance (do not edit)", "provenance", "json"),
]
TRUST: list[Column] = [
    ("Signal ID", "id", "id"),
    ("Dependency IDs", "dependencies", "list"),
    ("Trust Signal Category", "category", ("enum", "trust_category")),
    ("Specific Trust Signal", "claim", "text"),
    ("What Decision/Claim Depends On It", "reliance", "text"),
    ("Consequence If Trust Signal Is False", "consequence_if_false", "text"),
    ("Verification Status", "verification.status", ("enum", "verification_status")),
    ("Verification Method", "verification.method", ("enum", "method")),
    ("Last Verified", "verification.last_verified", "text"),
    ("Verification Currency", "verification.currency", "text"),
    ("Scope Match? (Y/N/Partial)", "verification.scope_match", ("enum", "scope_match")),
    ("Verification Would Require: Cost", "verification_requirements.cost", "text"),
    ("Verification Would Require: Access", "verification_requirements.access", "text"),
    ("Verification Would Require: Expertise", "verification_requirements.expertise", "text"),
    ("Verification Would Require: Time", "verification_requirements.time", "text"),
    ("Notes", "notes", "text"),
    ("Provenance (do not edit)", "provenance", "json"),
]
CHAINS: list[Column] = [
    ("Signal ID", "signal", "text"),
    ("Layer", "layer", "text"),
    ("Party", "party", "text"),
    ("Which Trusts", "trusts", "text"),
    ("Verified At This Layer? (Y/N)", "verified", "bool"),
]
SUMMARY: list[Column] = [
    ("Dependency ID", "dependency", "text"),
    ("What Converges Here", "what_converges", "text"),
    ("Failure Modes", "failure_modes", "list"),
    ("Trust Signals", "trust_signals", "list"),
    ("Why Governance Missed It", "why_governance_missed", "text"),
    ("Regulatory Exposure", "regulatory_exposure", "text"),
    ("Regulatory Clauses", "clauses", "list"),
    ("Recommended Action", "recommended_action", "text"),
    ("Actions (See Action Catalogue)", "actions", "list"),
    ("Why No Catalogue Action Applies", "action_gap", "text"),
    ("Internal Document That Should Address It", "internal_document", "text"),
    ("Governance Change Required", "governance_change", "text"),
    ("Provenance (do not edit)", "provenance", "json"),
]
TIES: list[Column] = [
    ("Order (Dependency IDs, first ranked first)", "order", "list"),
    ("Reason", "reason", "text"),
    ("Provenance (do not edit)", "provenance", "json"),
]
BOUNDARY = [
    ("AI System Name", "name", "Official name/identifier of the AI system"),
    ("System Owner", "owner", "Business unit and individual responsible"),
    ("Regulatory Classification", "regulatory_classification", "EU AI Act risk tier / DORA criticality / NIS2 / Sector-specific"),
    ("System Boundary", "boundary", "What is in scope (Phase 1, Step 1.1)"),
    ("Primary Function", "primary_function", "What the system does in business terms"),
    ("Users / Consumers", "users", "Who uses the output (internal teams, customers, automated systems)"),
    ("Deployment Date", "deployment_date", "When the system went into production"),
    ("Last Material Change", "last_material_change", "Date and nature of last significant update"),
    ("Downstream Dependencies", "downstream_dependencies", "What decisions, processes, or systems depend on this AI output"),
    ("Upstream Data Sources", "upstream_data_sources", "All data sources feeding the system"),
    ("Human-in-the-Loop Design", "human_in_the_loop_design", "Is a human expected to review/override AI outputs? Describe the mechanism and any evidence of actual override"),
    ("AI Agents and Tools", "agents_and_tools", "Agents in or acting for the system, the tools and MCP servers they can call, the credentials they use"),
    ("Compliance Evidence Location", "compliance_evidence_location", "Where governance docs, risk assessments, certifications are stored"),
]
ENTERED_FACTORS = ("regulatory_exposure", "detection_deficit", "trust_depth", "blast_radius", "remediation_complexity")


# -- helpers ----------------------------------------------------------------------


def _get(entry: dict, path: str) -> Any:
    node: Any = entry
    for key in path.split("."):
        if not isinstance(node, dict):
            return None
        node = node.get(key)
    return node


def _set(entry: dict, path: str, value: Any) -> None:
    keys = path.split(".")
    node = entry
    for key in keys[:-1]:
        node = node.setdefault(key, {})
    node[keys[-1]] = value


class Vocabulary:
    """Labels for enumerated values, both ways."""

    _SOURCES = {
        "layer": ("taxonomy/layers.yaml", ("layers",)),
        "owner_type": ("taxonomy/substrate.yaml", ("owner_types", "values")),
        "visibility": ("taxonomy/substrate.yaml", ("visibility", "values")),
        "fallback": ("taxonomy/substrate.yaml", ("fallback", "values")),
        "type": ("taxonomy/failure.yaml", ("failure_types", "values")),
        "latency": ("taxonomy/failure.yaml", ("detection_latency", "values")),
        "confidence": ("taxonomy/failure.yaml", ("detection_confidence", "values")),
        "impact": ("taxonomy/failure.yaml", ("impact", "values")),
        "materialisation_horizon": ("taxonomy/failure.yaml", ("materialisation_horizon", "values")),
        "trust_category": ("taxonomy/trust.yaml", ("categories", "values")),
        "verification_status": ("taxonomy/trust.yaml", ("verification_status", "values")),
        "method": ("taxonomy/trust.yaml", ("verification_method", "values")),
        "scope_match": ("taxonomy/trust.yaml", ("scope_match", "values")),
        "severity": ("rules/severity.yaml", ("levels",)),
        "category": ("rules/category.yaml", ("categories",)),
    }

    def __init__(self, pack: MethodPack):
        self.values: dict[str, list[dict]] = {}
        for name, (doc, keys) in self._SOURCES.items():
            node: Any = pack.doc(doc)
            for key in keys:
                node = node[key]
            self.values[name] = node

    def labels(self, name: str) -> list[str]:
        return [v.get("workbook_label", v["name"]) for v in self.values[name]]

    def label(self, name: str, value: Any) -> str | None:
        for v in self.values[name]:
            if v["id"] == value:
                return v.get("workbook_label", v["name"])
        return None if value is None else str(value)

    def value(self, name: str, text: Any) -> Any:
        if text is None or str(text).strip() == "":
            return None
        wanted = " ".join(str(text).split()).casefold()
        for v in self.values[name]:
            if wanted in {str(v["id"]).casefold(), v["name"].casefold(), str(v.get("workbook_label", "")).casefold()}:
                return v["id"]
        return str(text).strip()  # left for validation to reject, with the rule and a valid example


def _cell_text(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _bool_in(value: Any) -> Any:
    value = _cell_text(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    return {"y": True, "yes": True, "true": True, "n": False, "no": False, "false": False}.get(text, str(value))


def _list_in(value: Any) -> list[str]:
    value = _cell_text(value)
    if value is None:
        return []
    return [part.strip() for part in re.split(r"[,;\n]", str(value)) if part.strip()]


def _out(kind: Any, value: Any, vocab: Vocabulary) -> Any:
    if value is None:
        return None
    if kind == "bool":
        return "Y" if value else "N"
    if kind == "list":
        return ", ".join(value) if value else None
    if kind == "json":
        return json_line(value)
    if isinstance(kind, tuple):
        return vocab.label(kind[1], value)
    return value


def _in(kind: Any, value: Any, vocab: Vocabulary) -> Any:
    if kind == "bool":
        return _bool_in(value)
    if kind == "list":
        return _list_in(value)
    if kind == "json":
        text = _cell_text(value)
        return json.loads(text) if text else None
    if isinstance(kind, tuple):
        return vocab.value(kind[1], _cell_text(value))
    value = _cell_text(value)
    return str(value) if value is not None and not isinstance(value, bool) else value


def _write_value(cell, value: Any) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # text that starts with '=' stays text, never a formula


# -- writing ------------------------------------------------------------------------


@dataclass
class _Writer:
    pack: MethodPack
    vocab: Vocabulary
    workbook: Any = None
    lists: Any = None
    list_columns: dict[str, str] = field(default_factory=dict)

    def new(self) -> None:
        self.workbook = openpyxl.Workbook()
        self.workbook.remove(self.workbook.active)
        props = self.workbook.properties
        props.creator = "OSRA-CODE"
        props.lastModifiedBy = "OSRA-CODE"
        props.created = _FIXED
        props.modified = _FIXED
        self.lists = None
        self.list_columns = {}

    def list_range(self, name: str) -> str:
        """A validation list on a hidden sheet, from the pack's labels."""
        if self.lists is None:
            self.lists = self.workbook.create_sheet("Lists")
            self.lists.sheet_state = "hidden"
        if name not in self.list_columns:
            column = get_column_letter(len(self.list_columns) + 1)
            labels = self.vocab.labels(name) if name not in ("bool",) else ["Y", "N"]
            self.lists[f"{column}1"] = name
            for i, label in enumerate(labels, start=2):
                _write_value(self.lists[f"{column}{i}"], label)
            self.list_columns[name] = f"Lists!${column}$2:${column}${len(labels) + 1}"
        return self.list_columns[name]

    def table(self, title: str, intro: str, columns: list[Column], entries: list[dict],
              computed: list[tuple[str, Callable[[dict], Any]]] = (), *, extra_rows: int = SPARE_ROWS) -> Any:
        ws = self.workbook.create_sheet(title)
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=13)
        _write_value(ws["A2"], intro)
        headers = [c[0] for c in columns] + [f"{name} (computed)" for name, _ in computed]
        for i, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=i, value=header)
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if i > len(columns):
                cell.fill = _COMPUTED_FILL
            ws.column_dimensions[get_column_letter(i)].width = 14 if i == 1 else 24
        last = 3 + len(entries) + extra_rows
        for r, entry in enumerate(entries, start=4):
            for c, (_, path, kind) in enumerate(columns, start=1):
                _write_value(ws.cell(row=r, column=c), _out(kind, _get(entry, path), self.vocab))
            for c, (_, compute) in enumerate(computed, start=len(columns) + 1):
                cell = ws.cell(row=r, column=c)
                _write_value(cell, compute(entry))
                cell.fill = _COMPUTED_FILL
        for c, (_, _, kind) in enumerate(columns, start=1):
            letter = get_column_letter(c)
            for r in range(4, last + 1):
                ws.cell(row=r, column=c).protection = Protection(locked=False)
            if kind == "bool" or isinstance(kind, tuple):
                source = self.list_range("bool" if kind == "bool" else kind[1])
                rule = DataValidation(type="list", formula1=source, allow_blank=True)
                rule.add(f"{letter}4:{letter}{last}")
                ws.add_data_validation(rule)
        ws.freeze_panes = "B4"
        ws.protection.sheet = True
        ws.protection.formatCells = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        return ws

    def guide(self, title: str, rows: list[list[Any]], headers: list[str]) -> None:
        ws = self.workbook.create_sheet(title)
        for i, header in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=header).font = _HEADER_FONT
            ws.column_dimensions[get_column_letter(i)].width = 28 if i == 1 else 70
        for r, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c)
                _write_value(cell, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.protection.sheet = True

    def register_sheet(self, kinds: list[tuple[str, dict | None]]) -> None:
        ws = self.workbook.create_sheet("Register State")
        headers = ["Register", "State", "Confirmed By", "Confirmed At", "Confirmed Through", "Interactive (Y/N)",
                   "Retired Identifiers"]
        for i, header in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=header).font = _HEADER_FONT
            ws.column_dimensions[get_column_letter(i)].width = 22
        for r, (kind, doc) in enumerate(kinds, start=2):
            doc = doc or {}
            confirmation = doc.get("confirmation") or {}
            values = [kind, doc.get("state", "draft"), confirmation.get("by"), confirmation.get("at"),
                      confirmation.get("surface"),
                      None if not confirmation else ("Y" if confirmation.get("interactive") else "N"),
                      ", ".join(doc.get("retired", [])) or None]
            for c, value in enumerate(values, start=1):
                _write_value(ws.cell(row=r, column=c), value)
        ws.protection.sheet = True

    def meta(self, workbook_kind: str, payload: dict) -> None:
        ws = self.workbook.create_sheet("OSRA")
        ws.sheet_state = "hidden"
        rows = {"format": FORMAT, "format_version": FORMAT_VERSION, "workbook": workbook_kind,
                "method_pack": f"{self.pack.id} {self.pack.version}", **payload}
        for r, (key, value) in enumerate(rows.items(), start=1):
            ws.cell(row=r, column=1, value=key)
            _write_value(ws.cell(row=r, column=2), value if isinstance(value, (str, int)) else json_line(value))
        ws.protection.sheet = True

    def save(self, path: Path) -> None:
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        path.write_bytes(_normalise_archive(buffer.getvalue()))


def _normalise_archive(data: bytes) -> bytes:
    """Fixed entry timestamps and document dates, so that the same content
    always produces the same file (TR-63)."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", content)
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            target.writestr(entry, content)
    return out.getvalue()


def _content_hash(doc: dict | None) -> str:
    """What the register held at export, so an import can tell whether it
    was edited in the workbook (a confirmed register edited there returns
    to draft, FR-03)."""
    body = {k: v for k, v in (doc or {}).items() if k not in ("state", "confirmation")}
    return "sha256:" + hashlib.sha256(json_line(body).encode()).hexdigest()


def export(documents: dict[str, Any], results: dict | None, pack: MethodPack, outdir: Path) -> list[Path]:
    """Write the four workbooks for an assessment. With ``results``, the
    computed columns carry the engine's values; without, they are empty.
    An empty ``documents`` gives the blank templates."""
    vocab = Vocabulary(pack)
    writer = _Writer(pack, vocab)
    outdir.mkdir(parents=True, exist_ok=True)
    fm_results = {f["id"]: f for f in (results or {}).get("failure_modes", [])}
    ts_results = {t["id"]: t for t in (results or {}).get("trust_signals", [])}
    matrix = {row["dependency"]: row for row in (results or {}).get("matrix", [])}
    findings = {f["dependency"]: f for f in (results or {}).get("findings", [])}
    horizon_map = pack.rule("scoring.horizon")["map"]
    reg = lambda kind: documents.get(kind) or {}
    written = []

    # Phase 1
    writer.new()
    assessment = documents.get("assessment") or {}
    system = assessment.get("system") or {}
    ws = writer.workbook.create_sheet("System Boundary")
    ws["A1"] = "OSRA PHASE 1: SUBSTRATE MAP (v1.2)"
    ws["A1"].font = Font(bold=True, size=13)
    for i, header in enumerate(["Field", "Description", "Value (Complete This)"], start=1):
        ws.cell(row=3, column=i, value=header).font = _HEADER_FONT
    for r, (label, key, description) in enumerate(BOUNDARY, start=4):
        ws.cell(row=r, column=1, value=label)
        _write_value(ws.cell(row=r, column=2), description)
        _write_value(ws.cell(row=r, column=3), system.get(key))
        ws.cell(row=r, column=3).protection = Protection(locked=False)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 30, 60, 60
    ws.protection.sheet = True
    deps = sorted(reg("substrate").get("dependencies", []), key=lambda d: id_number(d["id"]))
    writer.table("Dependency Chain",
                 "One row per dependency. Every dependency keeps its ID for the whole assessment; Phases 2, 3 and 4 "
                 "refer to it by ID. Add rows below with the next free ID; never renumber or reuse an ID.",
                 SUBSTRATE, deps,
                 [("Category", lambda d: vocab.label("category", matrix.get(d["id"], {}).get("category")))])
    layers = pack.taxonomy("layers")["layers"]
    writer.guide("Layer Guide", [[layer.get("workbook_label", layer["name"]), " ".join(layer["questions"])] for layer in layers],
                 ["Dependency Layer", "Key Questions (Phase 1, Step 1.2)"])
    writer.register_sheet([("substrate", documents.get("substrate"))])
    meta = {"substrate": _content_hash(documents.get("substrate"))}
    if assessment:
        meta["assessment"] = {k: v for k, v in assessment.items() if k != "system"}
    writer.meta("phase1", meta)
    writer.save(outdir / FILES["phase1"])
    written.append(outdir / FILES["phase1"])

    # Phase 2
    writer.new()
    fms = sorted(reg("failures").get("failure_modes", []), key=lambda f: id_number(f["id"]))
    writer.table("Failure Surface Register",
                 "One row per failure mode. Severity comes from impact and a tested fallback only; the Silent Failure "
                 "Risk flag is a Silent failure with detection confidence Low or None. Both are computed by OSRA-CODE.",
                 FAILURES, fms,
                 [("Severity", lambda f: vocab.label("severity", fm_results.get(f["id"], {}).get("severity"))),
                  ("Silent Failure Risk", lambda f: _out("bool", fm_results.get(f["id"], {}).get("silent_failure_risk"), vocab)),
                  ("Horizon Score", lambda f: horizon_map.get(f.get("materialisation_horizon")) if f["id"] in fm_results else None)])
    failure = pack.taxonomy("failure")
    writer.guide("Failure Type Taxonomy",
                 [[v["workbook_label"], v["definition"], v["example"]] for v in failure["failure_types"]["values"]],
                 ["Failure Type", "Definition", "Example"])
    writer.guide("Severity Guide", [[v["name"], v["meaning"]] for v in failure["impact"]["values"]]
                 + [["Tested fallback", pack.rule("severity.tested-fallback")["note"]]],
                 ["Impact Level", "The failure affects..."])
    writer.guide("Detection-Horizon Guide",
                 [[f"Latency: {v['name']}", v["meaning"]] for v in failure["detection_latency"]["values"]]
                 + [[f"Confidence: {v['name']}", v["meaning"]] for v in failure["detection_confidence"]["values"]]
                 + [[f"Horizon: {v['workbook_label']}", v["meaning"]] for v in failure["materialisation_horizon"]["values"]],
                 ["Value", "Meaning"])
    writer.register_sheet([("failures", documents.get("failures"))])
    writer.meta("phase2", {"failures": _content_hash(documents.get("failures"))})
    writer.save(outdir / FILES["phase2"])
    written.append(outdir / FILES["phase2"])

    # Phase 3
    writer.new()
    signals = sorted(reg("trust").get("trust_signals", []), key=lambda t: id_number(t["id"]))
    writer.table("Trust Surface Register",
                 "One row per trust signal. List the dependency IDs it concerns, separated by commas. The trust gap and "
                 "chain depth are computed by OSRA-CODE.",
                 TRUST, signals,
                 [("Trust Gap", lambda t: _out("bool", ts_results.get(t["id"], {}).get("trust_gap"), vocab)),
                  ("Chain Depth", lambda t: ts_results.get(t["id"], {}).get("chain_depth"))])
    chains = [{"signal": t["id"], "layer": i, **layer} for t in signals for i, layer in enumerate(t.get("chain") or [], start=1)]
    writer.table("Trust Chains", "One row per layer of trust behind a signal, nearest layer first.", CHAINS, chains)
    trust = pack.taxonomy("trust")
    writer.guide("Verification Guide",
                 [[v["name"], v["definition"]] for v in trust["verification_status"]["values"]]
                 + [["Trust gap", pack.rule("trust-gap.gap")["note"]]],
                 ["Status", "Definition"])
    writer.register_sheet([("trust", documents.get("trust"))])
    writer.meta("phase3", {"trust": _content_hash(documents.get("trust"))})
    writer.save(outdir / FILES["phase3"])
    written.append(outdir / FILES["phase3"])

    # Phase 4
    writer.new()
    ws = writer.workbook.create_sheet("Convergence Matrix")
    ws["A1"] = "Convergence Matrix (computed by OSRA-CODE from Phases 1 to 3)"
    ws["A1"].font = Font(bold=True, size=13)
    headers = ["Dependency ID", "Cond. 1: Severity Critical or High", "Cond. 2: Silent Failure Risk", "Cond. 3: Trust Gap",
               "Single Point of Dependency", "Conditions Met", "Category", "Concentration Flag", "Clock"]
    for i, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=i, value=f"{header} (computed)")
        cell.font, cell.fill = _HEADER_FONT, _COMPUTED_FILL
        ws.column_dimensions[get_column_letter(i)].width = 20
    for r, row in enumerate((results or {}).get("matrix", []), start=4):
        values = [row["dependency"]] + [_out("bool", row[k], vocab) for k in ("condition_1", "condition_2", "condition_3", "single_point")]
        values += [row["conditions_met"], vocab.label("category", row["category"]), _out("bool", row["concentration_flag"], vocab),
                   ", ".join(row["clocks"]) or "Standard cycle"]
        for c, value in enumerate(values, start=1):
            _write_value(ws.cell(row=r, column=c), value)
    ws.protection.sheet = True
    names = {f["id"]: f["name"] for f in pack.rule_file("scoring")["factors"]}
    scoring_columns: list[Column] = [("Dependency ID", "dependency", "text")]
    for factor in ENTERED_FACTORS:
        scoring_columns += [(f"{names[factor]} (1-5)", f"factors.{factor}.score", "text"),
                            (f"{names[factor]}: Between Anchors? (Y/N)", f"factors.{factor}.between_anchors", "bool"),
                            (f"{names[factor]}: Reason", f"factors.{factor}.reason", "text")]
    scoring_columns.append(("Provenance (do not edit)", "provenance", "json"))
    scores = sorted(reg("scoring").get("scores", []), key=lambda s: id_number(s["dependency"]))
    writer.table("Convergence Scoring",
                 "Score every Critical Convergence, Convergence Point and Concentration Risk on the five factors a "
                 "person scores; where a finding sits between two anchors, take the lower score and give the reason. "
                 "Materialisation horizon, score and rank are computed.",
                 scoring_columns, scores,
                 [("Category", lambda s: vocab.label("category", matrix.get(s["dependency"], {}).get("category"))),
                  ("Materialisation Horizon", lambda s: findings.get(s["dependency"], {}).get("factors", {}).get("materialisation_horizon")),
                  ("Score", lambda s: findings.get(s["dependency"], {}).get("score")),
                  ("Rank", lambda s: findings.get(s["dependency"], {}).get("rank"))])
    writer.table("Tie Resolutions", "Findings the published tie-break cannot separate, in the order the practitioner "
                 "decided, with the reason.", TIES, reg("scoring").get("tie_resolutions", []), extra_rows=20)
    entries = sorted(reg("summary").get("entries", []), key=lambda e: id_number(e["dependency"]))
    writer.table("Convergence Risk Summary", "Per finding: what converges, why governance missed it, the regulatory "
                 "exposure and the recommended action (Phase 4, Step 4.3), and the governance it connects to (Step 4.4).",
                 SUMMARY, entries)
    scoring_doc = pack.rule_file("scoring")
    writer.guide("Scoring Anchors",
                 [[f"{f['name']} (x{f['weight']})", " ".join(f"{n} = {t}" for n, t in sorted(f["anchors"].items()))]
                  for f in scoring_doc["factors"]] + [["Lower anchor rule", scoring_doc["between_anchors"]["note"]]],
                 ["Factor", "Anchors"])
    writer.register_sheet([("scoring", documents.get("scoring")), ("summary", documents.get("summary"))])
    writer.meta("phase4", {"scoring": _content_hash(documents.get("scoring")),
                           "summary": _content_hash(documents.get("summary"))})
    writer.save(outdir / FILES["phase4"])
    written.append(outdir / FILES["phase4"])
    return written


# -- reading -------------------------------------------------------------------------


def _check_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        total = sum(info.file_size for info in archive.infolist())
        for info in archive.infolist():
            name = info.filename
            if name.startswith("/") or ".." in name.split("/") or "\\" in name:
                raise ValueError(f"unsafe path in archive: {name}")
        if total > MAX_ARCHIVE_BYTES:
            raise ValueError(f"the workbook expands to {total} bytes, over the limit of {MAX_ARCHIVE_BYTES}")


def _open(path: Path):
    _check_archive(path)
    return openpyxl.load_workbook(path, data_only=True, keep_links=False, keep_vba=False)


def _key(header: str) -> str:
    """A header's identity: its words before any parenthesis, so that the
    published headings ('Visibility (Visible / Known-Unmonitored /
    Invisible)') and the revised ones ('Visibility') match."""
    return re.sub(r"[^a-z0-9]+", " ", header.casefold().split("(")[0]).strip()


def _header_map(ws, header_row: int) -> dict[str, int]:
    found: dict[str, int] = {}
    for cell in ws[header_row]:
        if isinstance(cell.value, str):
            found.setdefault(_key(cell.value), cell.column)
    return found


def _find_header_row(ws, first: str) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=10):
        for cell in row:
            if isinstance(cell.value, str) and _key(cell.value) == _key(first):
                return cell.row
    return None


def _rows(ws, columns: list[Column], vocab: Vocabulary) -> list[dict]:
    header_row = _find_header_row(ws, columns[0][0])
    if header_row is None:
        return []
    headers = _header_map(ws, header_row)
    entries = []
    for row in ws.iter_rows(min_row=header_row + 1):
        values = {}
        empty = True
        for header, path, kind in columns:
            column = headers.get(_key(header))
            raw = row[column - 1].value if column and column - 1 < len(row) else None
            if _cell_text(raw) is not None:
                empty = False
            values[path] = _in(kind, raw, vocab)
        filled = [path for path, value in values.items() if value not in (None, [])]
        if empty or filled == ["id"]:
            continue  # an empty row, or one the template numbered in advance
        entry: dict[str, Any] = {}
        for path, value in values.items():
            if value is None or value == []:
                continue
            _set(entry, path, value)
        for path in _REQUIRED_NULLABLE.get(columns[0][0], ()):
            if _get(entry, path) is None:
                _set(entry, path, None)
        entries.append(entry)
    return entries


# Keys the schemas require to be present even when their value is null.
_REQUIRED_NULLABLE = {
    "Failure ID": ("detection.mechanism",),
    "Signal ID": ("reliance",),
}
# What a published-workbook row needs before it is imported; a row without it
# is a starter or an unfinished row, and is reported rather than imported.
_PUBLISHED_REQUIRED = {
    "substrate": ("id", "name", "layer", "single_point", "visibility", "fallback"),
    "failures": ("id", "dependency", "type", "description", "detection.latency", "detection.confidence", "impact",
                 "tested_fallback", "materialisation_horizon"),
    "trust": ("id", "dependencies", "category", "claim", "verification.status"),
}


def _complete(kind: str, rows: list[dict], notes: list, file: str) -> list[dict]:
    kept = []
    for row in rows:
        missing = [path for path in _PUBLISHED_REQUIRED[kind] if _get(row, path) in (None, [])]
        if kind == "substrate" and row.get("fallback") in ("yes", "partial") and row.get("fallback_tested") is None:
            missing.append("fallback_tested")
        if missing:
            notes.append(_note(file, row.get("id") or kind, f"not imported: the row has no {', '.join(missing)}"))
            continue
        if kind == "substrate" and "owner_type" not in row:
            row["owner_type"] = "unknown"
        kept.append(row)
    if kind == "substrate" and kept:
        notes.append(_note(file, "substrate", "the published workbook has no owner type; imported as unknown"))
    return kept


def _meta(wb) -> dict[str, Any] | None:
    if "OSRA" not in wb.sheetnames:
        return None
    ws = wb["OSRA"]
    meta = {}
    for row in ws.iter_rows(min_row=1, max_col=2):
        key, value = row[0].value, row[1].value
        if key:
            meta[key] = value
    return meta if meta.get("format") == FORMAT else None


def _registers_state(wb) -> dict[str, dict]:
    if "Register State" not in wb.sheetnames:
        return {}
    states = {}
    ws = wb["Register State"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        kind, state, by, at, surface, interactive, retired = (list(row) + [None] * 7)[:7]
        entry: dict[str, Any] = {"state": state or "draft"}
        if by:
            entry["confirmation"] = {"by": by, "at": at, "surface": surface, "interactive": interactive == "Y"}
        entry["retired"] = _list_in(retired)
        states[kind] = entry
    return states


@dataclass
class Imported:
    documents: dict[str, Any]
    notes: list[Diagnostic]
    revised: bool


def import_workbooks(paths: list[Path], pack: MethodPack) -> Imported:
    """Read the workbooks into an assessment's documents. Revised workbooks
    give back exactly what was exported; published v1.2 workbooks give a
    draft assessment and a note for everything they could not supply."""
    vocab = Vocabulary(pack)
    documents: dict[str, Any] = {}
    notes: list[Diagnostic] = []
    revised = True
    for path in paths:
        try:
            wb = _open(path)
        except (ValueError, zipfile.BadZipFile, OSError) as exc:
            notes.append(diagnostic("OSRA-E302", entity=path.name, file=path.name, detail=str(exc)))
            continue
        meta = _meta(wb)
        revised = revised and meta is not None
        states = _registers_state(wb) if meta else {}
        if meta is not None and meta.get("method_pack") != f"{pack.id} {pack.version}":
            notes.append(diagnostic("OSRA-E205", entity=path.name, field="method_pack", value=meta.get("method_pack"),
                                    expected=f"{pack.id} {pack.version}"))
        if "System Boundary" in wb.sheetnames:
            documents["assessment"] = _assessment(wb["System Boundary"], meta, pack, notes, path.name)
        if "Dependency Chain" in wb.sheetnames:
            rows = _rows(wb["Dependency Chain"], SUBSTRATE, vocab)
            if meta is None:
                rows = _complete("substrate", rows, notes, path.name)
            documents["substrate"] = _register("substrate", "dependencies", rows, states, meta, notes, path.name)
        if "Failure Surface Register" in wb.sheetnames:
            rows = _rows(wb["Failure Surface Register"], FAILURES, vocab)
            if meta is None:
                rows = _published_failures(wb["Failure Surface Register"], vocab, notes, path.name)
            documents["failures"] = _register("failures", "failure_modes", rows, states, meta, notes, path.name)
        if "Trust Surface Register" in wb.sheetnames:
            rows = _rows(wb["Trust Surface Register"], TRUST, vocab)
            if meta is None:
                rows = _published_trust(wb, vocab, notes, path.name)
            else:
                chains: dict[str, list] = {}
                for layer in _rows(wb["Trust Chains"], CHAINS, vocab) if "Trust Chains" in wb.sheetnames else []:
                    item = {"party": layer.get("party"), "verified": layer.get("verified")}
                    if layer.get("trusts") is not None:
                        item["trusts"] = layer["trusts"]
                    chains.setdefault(layer.get("signal"), []).append((int(layer.get("layer") or 0), item))
                for ts in rows:
                    if ts.get("id") in chains:
                        ts["chain"] = [item for _, item in sorted(chains[ts["id"]], key=lambda x: x[0])]
            documents["trust"] = _register("trust", "trust_signals", rows, states, meta, notes, path.name)
        if "Convergence Scoring" in wb.sheetnames:
            documents["scoring"] = _scoring(wb, vocab, states, meta, notes, path.name, pack)
        if "Convergence Risk Summary" in wb.sheetnames and meta is not None:
            documents["summary"] = _register("summary", "entries", _rows(wb["Convergence Risk Summary"], SUMMARY, vocab),
                                             states, meta, notes, path.name)
        elif "Convergence Risk Summary" in wb.sheetnames:
            notes.append(_note(path.name, "Convergence Risk Summary",
                               "the narrative is keyed by rank in the published workbook and is not imported; "
                               "write it with 'osra-code summary' after scoring"))
    return Imported(documents=documents, notes=notes, revised=revised)


def _note(file: str, entity: str, detail: str) -> Diagnostic:
    return diagnostic("OSRA-W622", entity=entity, file=file, detail=detail)


def _assessment(ws, meta, pack, notes, file) -> dict:
    system = {}
    labels = {" ".join(label.split()).casefold(): key for label, key, _ in BOUNDARY}
    labels["ai system name"] = "name"
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or not isinstance(row[0], str):
            continue
        key = labels.get(" ".join(row[0].split()).casefold())
        value = _cell_text(row[2]) if len(row) > 2 else None
        if key and value is not None:
            system[key] = str(value)
    base = json.loads(meta["assessment"]) if meta and meta.get("assessment") else None
    if base is None:
        base = {"schema_version": 1, "kind": "assessment", "method_pack": {"id": pack.id, "version": pack.version},
                "assessment_type": "full", "deployment_mode": {"declared": "C"}, "agent_access": "read-only",
                "status": "active"}
        if "boundary" not in system:
            notes.append(_note(file, "System Boundary", "the published workbook has no System Boundary field; "
                                                        "give it with --boundary"))
        notes.append(_note(file, "assessment.yaml", "assessment type, declared deployment mode and agent access are not "
                                                    "in the published workbook; set to full, C and read-only"))
    return {**base, "system": system}


def _register(kind, list_key, rows, states, meta, notes, file, extra: dict | None = None) -> dict:
    state = states.get(kind, {})
    doc: dict[str, Any] = {"schema_version": 1, "kind": kind, "state": state.get("state", "draft")}
    if state.get("confirmation"):
        doc["confirmation"] = state["confirmation"]
    if state.get("retired"):
        doc["retired"] = state["retired"]
    doc[list_key] = rows
    doc.update(extra or {})
    if meta is None:
        doc["state"] = "draft"
        notes.append(_note(file, kind, "the published workbook carries no register state or provenance; imported as "
                                       "draft, attributed to the import"))
    elif doc["state"] == "confirmed" and meta.get(kind) != _content_hash(doc):
        doc["state"] = "draft"
        doc.pop("confirmation", None)
        notes.append(_note(file, kind, "the register was edited in the workbook after it was confirmed; it returns to "
                                       "draft (FR-03)"))
    return doc


def _scoring(wb, vocab, states, meta, notes, file, pack) -> dict:
    names = {f["id"]: f["name"] for f in pack.rule_file("scoring")["factors"]}
    if meta is not None:
        columns: list[Column] = [("Dependency ID", "dependency", "text")]
        for factor in ENTERED_FACTORS:
            columns += [(f"{names[factor]} (1-5)", f"factors.{factor}.score", "text"),
                        (f"{names[factor]}: Between Anchors? (Y/N)", f"factors.{factor}.between_anchors", "bool"),
                        (f"{names[factor]}: Reason", f"factors.{factor}.reason", "text")]
        columns.append(("Provenance (do not edit)", "provenance", "json"))
        scores = _rows(wb["Convergence Scoring"], columns, vocab)
        ties = _rows(wb["Tie Resolutions"], TIES, vocab) if "Tie Resolutions" in wb.sheetnames else []
    else:
        published = [("Dependency ID", "dependency", "text")] + [
            (f"{names[f].title()} (1-5)", f"factors.{f}.score", "text") for f in ENTERED_FACTORS]
        scores = []
        for row in _rows(wb["Convergence Scoring"], published, vocab):
            if re.fullmatch(r"DEP-[0-9]{2,}", str(row.get("dependency", ""))):
                if row.get("factors"):
                    scores.append(row)
            elif any(isinstance(v.get("score"), int) or str(v.get("score", "")).isdigit()
                     for v in (row.get("factors") or {}).values()):
                notes.append(_note(file, str(row.get("dependency")), "not imported: the row's Dependency ID is not "
                                                                     "an identifier such as DEP-01"))
            # rows with text and no scores are the sheet's guidance, not data
        ties = []
        notes.append(_note(file, "scoring", "factor reasons, between-anchor marks and tie resolutions are not in the "
                                            "published workbook; materialisation horizon, score and rank are recomputed"))
    for entry in scores:
        for factor in list(entry.get("factors", {})):
            value = entry["factors"][factor].get("score")
            if isinstance(value, str) and value.isdigit():
                entry["factors"][factor]["score"] = int(value)
    return _register("scoring", "scores", [s for s in scores if s.get("dependency")], states, meta, notes, file,
                     {"tie_resolutions": ties} if ties else None)


def _published_failures(ws, vocab, notes, file) -> list[dict]:
    """The published v1.2 Failure Surface Register: one propagation column,
    and computed columns that are ignored and recomputed (TR-60)."""
    columns = [c for c in FAILURES if c[1] not in ("propagation.path", "propagation.downstream_impact",
                                                    "propagation.downstream_detection", "notes", "provenance")]
    columns = [("Tested Fallback In Place? (Y/N)", "tested_fallback", "bool") if c[1] == "tested_fallback" else c for c in columns]
    columns.append(("Downstream Impact & Propagation Path", "propagation.downstream_impact", "text"))
    columns.append(("Severity", "_severity", "text"))
    rows = _complete("failures", _rows(ws, columns, vocab), notes, file)
    for row in rows:
        row.pop("_severity", None)
        row.setdefault("detection", {}).setdefault("mechanism", None)
    if rows:
        notes.append(_note(file, "failures", "propagation is one column in the published workbook and is imported as "
                                             "the downstream impact; severity and the silent failure flag are recomputed"))
    return rows


def _published_trust(wb, vocab, notes, file) -> list[dict]:
    columns = [c for c in TRUST if c[1] in ("id", "dependencies", "category", "claim", "reliance",
                                            "consequence_if_false", "verification.status", "verification.method",
                                            "verification.last_verified", "verification.scope_match")]
    columns = [("Dependency ID(s) (from Phase 1)", "dependencies", "list") if c[1] == "dependencies" else c for c in columns]
    columns = [("Scope Match? (Y/N/Partial)", "verification.scope_match", ("enum", "scope_match"))
               if c[1] == "verification.scope_match" else c for c in columns]
    columns = [("Consequence If Trust Signal Is False", "consequence_if_false", "text") if c[1] == "consequence_if_false" else c
               for c in columns]
    rows = _complete("trust", _rows(wb["Trust Surface Register"], columns, vocab), notes, file)
    for row in rows:
        row.setdefault("reliance", None)
    if "Trust Chains" in wb.sheetnames:
        ws = wb["Trust Chains"]
        header_row = _find_header_row(ws, "Signal ID")
        by_id = {r.get("id"): r for r in rows}
        for row in ws.iter_rows(min_row=(header_row or 1) + 1, values_only=True):
            if not row or not row[0] or row[0] not in by_id:
                continue
            parties = [_cell_text(v) for v in row[2:5]]
            verified = _bool_in(row[6]) if len(row) > 6 else None
            layers = [p for p in parties if p is not None]
            if layers:
                by_id[row[0]]["chain"] = [{"party": str(p), "verified": verified is True} for p in layers]
        notes.append(_note(file, "trust chains", "the published workbook records one 'verified at each level' answer "
                                                 "per chain; it is applied to every layer"))
    return rows
